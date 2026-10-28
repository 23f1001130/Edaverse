"""Excel-to-CSV iteration, with an opt-in native reader for larger servers."""
import csv
import datetime as dt
import io
import logging
import math
import os
import posixpath
import xml.etree.ElementTree as ET
import xml.parsers.expat as expat
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

logger = logging.getLogger(__name__)
# Native readers materialize a sheet. Keep unusually large archives on the
# read-only path, even on the 2 GiB deployment. These are guards, not RAM promises.
NATIVE_MAX_FILE_BYTES = 64 * 1024 * 1024
NATIVE_MAX_EXPANDED_BYTES = 384 * 1024 * 1024
NATIVE_MAX_CELLS = 8_000_000


class _UseReadOnlyReader(Exception):
    pass


def _inspect_native_sheet(source):
    """Find the first worksheet and reject features the native reader loses.

    Calamine returns Excel error cells as empty strings. Inspect XML without
    constructing a tree for the sheet, so those files use openpyxl instead.
    Dimensions preserve leading columns and trailing empty/formatted cells.
    """
    source.seek(0)
    with zipfile.ZipFile(source) as archive:
        if sum(info.file_size for info in archive.infolist()) > NATIVE_MAX_EXPANDED_BYTES:
            raise _UseReadOnlyReader()
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        worksheets = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels
                      if rel.attrib.get("Type", "").endswith("/worksheet")}
        sheets = [sheet for sheet in workbook.iter() if sheet.tag.endswith("}sheet")
                  and any(value in worksheets for key, value in sheet.attrib.items()
                          if key.endswith("}id"))]
        if not sheets:
            raise _UseReadOnlyReader()
        relation = next(value for key, value in sheets[0].attrib.items() if key.endswith("}id"))
        target = worksheets[relation]
        path = posixpath.normpath(target.lstrip("/") if target.startswith("/")
                                 else posixpath.join("xl", target))
        bounds = None
        cells = 0
        max_row = 0
        max_column = ""

        def start_element(name, attrs):
            nonlocal bounds, cells, max_row, max_column
            tag = name.rsplit("}", 1)[-1]
            if tag == "dimension" and attrs.get("ref"):
                bounds = range_boundaries(attrs["ref"])
            if tag == "c":
                cells += 1
                if attrs.get("t") == "e" or cells > NATIVE_MAX_CELLS:
                    raise _UseReadOnlyReader()
                reference = attrs.get("r", "")
                column = reference.rstrip("0123456789")
                if not column or column == reference:
                    raise _UseReadOnlyReader()
                max_row = max(max_row, int(reference[len(column):]))
                if (len(column), column) > (len(max_column), max_column):
                    max_column = column

        xml = expat.ParserCreate(namespace_separator="}")
        xml.StartElementHandler = start_element
        # XLSX worksheets don't need DTDs or entities.
        def reject_doctype(*args):
            raise _UseReadOnlyReader()

        xml.StartDoctypeDeclHandler = reject_doctype
        with archive.open(path) as sheet:
            while chunk := sheet.read(1024 * 1024):
                xml.Parse(chunk, False)
            xml.Parse(b"", True)
        if bounds and bounds[2] * bounds[3] > NATIVE_MAX_CELLS:
            raise _UseReadOnlyReader()
        # A sparse sheet can declare tiny dimensions yet force the native reader
        # to allocate a very large rectangle. Check actual cell coordinates too.
        if max_column and range_boundaries(f"{max_column}{max_row}")[2] * max_row > NATIVE_MAX_CELLS:
            raise _UseReadOnlyReader()
    return sheets[0].attrib["name"], len(sheets), bounds


def _native_value(value):
    # Match openpyxl's values before csv.writer serializes them. In particular,
    # retain text IDs such as '001', while writing integral numeric cells as '1'.
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return int(value)
    if isinstance(value, dt.date) and not isinstance(value, dt.datetime):
        return dt.datetime.combine(value, dt.time())
    return value


def _write_rows(rows, output, header_row):
    writer = csv.writer(output)
    preview = []
    count = 0
    for index, row in enumerate(rows):
        if index < 8:
            preview.append(row)
        if index >= (header_row or 0):
            writer.writerow(row)
            count += 1
    preview_text = io.StringIO()
    csv.writer(preview_text).writerows(preview)
    return max(0, count - 1), preview_text.getvalue().encode("utf-8")


def write_xlsx_csv(source, size, output, header_row=None):
    """Return (data rows, preview bytes, sheet count, reader name).

    Default to openpyxl for small-memory hosts. The Azure image opts in through
    XLSX_READER=calamine; an unsupported file restarts with the read-only reader.
    """
    if os.environ.get("XLSX_READER", "openpyxl").lower() == "calamine" and size <= NATIVE_MAX_FILE_BYTES:
        try:
            from python_calamine import CalamineWorkbook
            name, sheet_count, bounds = _inspect_native_sheet(source)
            source.seek(0)
            with CalamineWorkbook.from_filelike(source) as workbook:
                sheet = workbook.get_sheet_by_name(name)
                left = sheet.start[1] if sheet.start else 0
                width = bounds[2] if bounds else left + sheet.width
                height = bounds[3] if bounds else (sheet.end[0] + 1 if sheet.end else 0)

                def rows():
                    count = 0
                    for row in sheet.iter_rows():
                        if count >= height:
                            break
                        values = [None] * left + [_native_value(value) for value in row]
                        yield values[:width] + [None] * max(0, width - len(values))
                        count += 1
                    for _ in range(count, height):
                        yield [None] * width

                total, preview = _write_rows(rows(), output, header_row)
            return total, preview, sheet_count, "calamine"
        except Exception:
            # Never append a second import after a partial native attempt.
            output.seek(0)
            output.truncate()
            logger.info("Using read-only Excel reader for this workbook")
    source.seek(0)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        total, preview = _write_rows(workbook.worksheets[0].iter_rows(values_only=True), output, header_row)
        return total, preview, len(workbook.worksheets), "openpyxl"
    finally:
        workbook.close()
